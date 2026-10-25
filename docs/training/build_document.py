#!/usr/bin/env python3
"""Build the Word manual from Markdown; dependencies: requirements.txt."""
from pathlib import Path
import datetime
import re
from markdown_it import MarkdownIt
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
STEM = '5G通信组网与运维培训'
MD = MarkdownIt('commonmark').enable('table')


def font(run, size=None, bold=False, color='24384C', mono=False):
    run.font.name = 'DejaVu Sans Mono' if mono else 'Noto Sans CJK SC'
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC')
    if size: run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def field(p, instruction):
    element=OxmlElement('w:fldSimple');element.set(qn('w:instr'),instruction)
    p._p.append(element)


def toc_field(p):
    # A complex TOC field is recognized as a document index by Word and Writer.
    begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
    begin.set(qn('w:dirty'),'true');p.add_run()._r.append(begin)
    instruction=OxmlElement('w:instrText');instruction.set(qn('xml:space'),'preserve')
    instruction.text=' TOC \\o "1-1" \\h \\z \\u '
    p.add_run()._r.append(instruction)
    separate=OxmlElement('w:fldChar');separate.set(qn('w:fldCharType'),'separate')
    p.add_run()._r.append(separate)
    p.add_run('更新目录后显示章节与页码')
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');p.add_run()._r.append(end)


def shade(p, fill='F0F4F8'):
    e=OxmlElement('w:shd');e.set(qn('w:fill'),fill)
    p._p.get_or_add_pPr().append(e)


def inline(p, tokens, size=10, bold=False, color='24384C'):
    strong=bold; italic=False; link=None
    for t in tokens or []:
        if t.type=='strong_open': strong=True
        elif t.type=='strong_close': strong=bold
        elif t.type=='em_open': italic=True
        elif t.type=='em_close': italic=False
        elif t.type=='link_open': link=t.attrGet('href')
        elif t.type=='link_close': link=None
        elif t.type=='image':
            image=ROOT/t.attrGet('src')
            with Image.open(image) as img: w,h=img.size
            picture=p.add_run().add_picture(str(image),width=Cm(min(16.8,18*w/h)))
            picture._inline.docPr.set('descr',t.content)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            caption=p._parent.add_paragraph()
            caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.keep_with_next=True
            font(caption.add_run(t.content),9,color='58708B')
        elif t.type in ('text','code_inline','softbreak','hardbreak'):
            r=p.add_run('\n' if t.type in ('softbreak','hardbreak') else t.content)
            font(r,size,strong,color,t.type=='code_inline');r.italic=italic
            if link:
                r.font.color.rgb=RGBColor.from_string('176A9B')
                rel=p.part.relate_to(link,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
                h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
                r._r.addprevious(h);h.append(r._r)


def build():
    doc=Document();section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(1.7);section.bottom_margin=Cm(1.7)
    section.left_margin=Cm(2.1);section.right_margin=Cm(2.1)
    section.header_distance=Cm(.7);section.footer_distance=Cm(.7)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
        s=doc.styles[name];s.font.name='Noto Sans CJK SC'
        s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC')
    normal=doc.styles['Normal'];normal.font.size=Pt(10)
    normal.paragraph_format.line_spacing=1.13
    normal.paragraph_format.space_after=Pt(4)
    for name,size in [('Title',28),('Heading 1',18),('Heading 2',13),('Heading 3',11)]:
        s=doc.styles[name];s.font.size=Pt(size);s.font.bold=True
        s.font.color.rgb=RGBColor.from_string('163E64')
        s.paragraph_format.space_before=Pt(11);s.paragraph_format.space_after=Pt(6)
        s.paragraph_format.keep_with_next=True
    doc.styles['Title'].paragraph_format.space_before=Pt(48)
    header=section.header.paragraphs[0]
    font(header.add_run('5G 通信组网与运维培训  ·  V1.0'),8,color='58708B')
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(footer.add_run('2026-10-25  |  '),8,color='58708B');field(footer,'PAGE')
    font(footer.add_run(' / '),8,color='58708B');field(footer,'NUMPAGES')
    doc.core_properties.title=STEM
    doc.core_properties.subject='半天入门课程与课后操作手册'
    doc.core_properties.author='XCN 培训资料'
    doc.core_properties.created=datetime.datetime(2026,10,25,tzinfo=datetime.timezone.utc)
    doc.core_properties.modified=doc.core_properties.created
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
    tokens=MD.parse((ROOT/(STEM+'.md')).read_text());i=0;lists=[];count=0;quote=False
    while i<len(tokens):
        t=tokens[i]
        if t.type=='heading_open':
            level=int(t.tag[1:]);p=doc.add_paragraph(style='Title' if level==1 else f'Heading {min(level-1,3)}')
            if level==2 or tokens[i+1].content in (
                '2.6 CNI、Multus 与 SR-IOV Device Plugin',
                '6.3 k9s 的启动和快捷键', 'A.3 常见术语速查',
                'C.3 模拟 UE：UspLbLinkVppMemifSessionCreate'):
                p.paragraph_format.page_break_before=True
            count+=1
            bookmark=OxmlElement('w:bookmarkStart');bookmark.set(qn('w:id'),str(count));bookmark.set(qn('w:name'),f'section_{count}');p._p.append(bookmark)
            inline(p,tokens[i+1].children,size={1:28,2:18,3:13}.get(level,11),bold=True,color='163E64')
            end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(count));p._p.append(end)
            if tokens[i+1].content=='阅读导航与课程安排':
                toc=doc.add_paragraph();toc_field(toc)
            i+=3;continue
        if t.type in ('bullet_list_open','ordered_list_open'):lists.append(t.type)
        elif t.type in ('bullet_list_close','ordered_list_close'):lists.pop()
        elif t.type=='blockquote_open':quote=True
        elif t.type=='blockquote_close':quote=False
        elif t.type=='paragraph_open':
            style=None
            if lists:style='List Number' if lists[-1]=='ordered_list_open' else 'List Bullet'
            p=doc.add_paragraph(style=style);inline(p,tokens[i+1].children)
            if quote:shade(p,'FFF1E5')
            i+=3;continue
        elif t.type=='fence':
            p=doc.add_paragraph();p.paragraph_format.line_spacing=1.0
            p.paragraph_format.keep_together=True
            p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(6)
            p.paragraph_format.left_indent=Cm(.12)
            shade(p);font(p.add_run(t.content.rstrip()),8,mono=True,color='203B56')
        elif t.type=='table_open':
            rows=[];row=[];i+=1
            while tokens[i].type!='table_close':
                tt=tokens[i]
                if tt.type=='tr_open':row=[]
                elif tt.type=='inline':row.append(tt)
                elif tt.type=='tr_close':rows.append(row)
                i+=1
            cols=len(rows[0]);table=doc.add_table(rows=0,cols=cols);table.style='Table Grid';table.autofit=False
            ratios={2:[.32,.68],3:[.25,.34,.41],4:[.22,.26,.25,.27],5:[.16,.22,.12,.24,.26]}.get(cols,[1/cols]*cols)
            for c,ratio in zip(table.columns,ratios):c.width=Cm(16.8*ratio)
            for ri,row in enumerate(rows):
                cells=table.add_row().cells
                for cell,content,ratio in zip(cells,row,ratios):
                    cell.width=Cm(16.8*ratio);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p=cell.paragraphs[0];p.paragraph_format.line_spacing=1.05
                    p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
                    inline(p,content.children,size=8.5,bold=ri==0,color='FFFFFF' if ri==0 else '24384C')
                    if ri==0 or ri%2==0:
                        e=OxmlElement('w:shd');e.set(qn('w:fill'),'214D72' if ri==0 else 'F4F7FA');cell._tc.get_or_add_tcPr().append(e)
                pr=cells[0]._tc.getparent().get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
                if ri==0:pr.append(OxmlElement('w:tblHeader'))
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=.2
        i+=1
    out=ROOT/(STEM+'.docx');doc.save(out);print(out)


if __name__=='__main__':build()
